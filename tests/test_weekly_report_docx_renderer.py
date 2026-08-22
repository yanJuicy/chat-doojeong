from __future__ import annotations

import unittest
from datetime import date
from io import BytesIO

from docx import Document as DocxDocument

from app.core.weekly_report_composer import ReportItem, ReportPeriodBlock, WeeklyReportView
from app.core.weekly_report_docx_renderer import render_weekly_report_docx


def _view(items: list[str], *, source_format: str | None) -> WeeklyReportView:
    period = (date(2026, 8, 17), date(2026, 8, 23))
    return WeeklyReportView(
        department="개발팀",
        current_week=ReportPeriodBlock(
            period_start=period[0],
            period_end=period[1],
            items=[ReportItem(id=str(i), content=text) for i, text in enumerate(items)],
        ),
        next_week=ReportPeriodBlock(period_start=period[0], period_end=period[1], items=[]),
        source_format=source_format,
    )


def _current_week_cell_text(docx_bytes: bytes) -> str:
    doc = DocxDocument(BytesIO(docx_bytes))
    table = doc.tables[0]
    return table.cell(1, 1).text


class RenderFormatTests(unittest.TestCase):
    def test_default_bullet_when_no_source_format(self) -> None:
        docx_bytes = render_weekly_report_docx(_view(["항목1", "항목2"], source_format=None))
        text = _current_week_cell_text(docx_bytes)
        self.assertIn("• 항목1", text)
        self.assertIn("• 항목2", text)

    def test_dash_marker_reproduced(self) -> None:
        docx_bytes = render_weekly_report_docx(_view(["항목1", "항목2"], source_format="bullet:-"))
        text = _current_week_cell_text(docx_bytes)
        self.assertIn("- 항목1", text)
        self.assertIn("- 항목2", text)
        self.assertNotIn("•", text)

    def test_prose_format_joins_items_without_bullets(self) -> None:
        docx_bytes = render_weekly_report_docx(_view(["항목1을 진행함", "항목2를 완료함"], source_format="prose"))
        text = _current_week_cell_text(docx_bytes)
        self.assertNotIn("•", text)
        self.assertNotIn("- ", text)
        self.assertIn("항목1을 진행함", text)
        self.assertIn("항목2를 완료함", text)

    def test_empty_items_still_render_dash_placeholder(self) -> None:
        docx_bytes = render_weekly_report_docx(_view([], source_format="bullet:-"))
        text = _current_week_cell_text(docx_bytes)
        self.assertEqual(text.strip(), "-")


if __name__ == "__main__":
    unittest.main()
